"""TRACE audit workflow API.

Run with: uvicorn backend.main:app --reload --port 8000
"""
import base64, csv, hashlib, hmac, io, json, os, secrets, time
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Annotated, Any, Literal
from urllib.parse import urlencode
from uuid import uuid4

from fastapi import Depends, FastAPI, File, HTTPException, Query, Request, UploadFile, status
from fastapi.concurrency import run_in_threadpool
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import HTMLResponse, RedirectResponse, Response
from jose import JWTError, jwt
from pydantic import BaseModel, ConfigDict, EmailStr, Field, field_validator
from sqlalchemy import select
from sqlalchemy.orm import Session
from dotenv import load_dotenv

load_dotenv()

from .database import Base, DATA_DIR, engine, get_db
from .models import AccessGrant, AdminAuditEvent, Artifact, AuditEvent, AuditSnapshot, Company, Engagement, EvidenceRequest, Finding, ProcedureRun, ReportVersion, SamplingRun, SecurityEvent, User, now
from .migrations import migrate
from .services.data_pipeline import DatasetError, SUPPORTED_EXTENSIONS, dataframe_records, load_dataframe, profile_dataframe
from .services.audit_engine import CHECKS, deterministic_sample, execute_check, risk_indicators, run_checks
from .services.revenue_analysis import analyze_revenue_datasets

Base.metadata.create_all(bind=engine)
migrate()
SECRET_KEY = os.getenv("TRACE_SECRET_KEY", "local-development-secret-change-before-deployment")
ALGORITHM = "HS256"
TOKEN_MINUTES = int(os.getenv("TRACE_TOKEN_MINUTES", "1440"))
GOOGLE_CLIENT_ID = os.getenv("TRACE_GOOGLE_CLIENT_ID", "").strip()
GOOGLE_CLIENT_SECRET = os.getenv("TRACE_GOOGLE_CLIENT_SECRET", "").strip()
GOOGLE_REDIRECT_URI = os.getenv("TRACE_GOOGLE_REDIRECT_URI", "http://127.0.0.1:8000/api/v1/auth/google/callback").strip()
PLATFORM_ADMIN_EMAIL = "drb.genai@gmail.com"
UPLOAD_DIR = DATA_DIR / "uploads"
UPLOAD_DIR.mkdir(exist_ok=True)
MAX_UPLOAD_BYTES = int(os.getenv("TRACE_MAX_UPLOAD_BYTES", str(50 * 1024 * 1024)))

app = FastAPI(title="TRACE — AI-Run Statutory Revenue Audit", version="0.2.0", description="Evidence-led, reproducible audit workflow API")
origins = [x.strip() for x in os.getenv("TRACE_CORS_ORIGINS", "http://localhost:5173,http://127.0.0.1:5173").split(",") if x.strip()]
app.add_middleware(CORSMiddleware, allow_origins=origins, allow_credentials=True, allow_methods=["*"], allow_headers=["Authorization", "Content-Type"])

def user_out(user: User):
    return {"id": user.id, "email": user.email, "name": user.name, "role": user.role, "organization": user.organization, "company_id": user.company_id, "active": user.active, "created_at": user.created_at}

def normalized_name(name: str) -> str:
    return " ".join(name.casefold().split()).strip()

def get_or_create_company(db: Session, name: str) -> Company:
    label = " ".join(name.split()).strip()
    if not label:
        raise HTTPException(422, "Company name is required")
    key = normalized_name(label)
    company = db.scalar(select(Company).where(Company.normalized_name == key))
    if company:
        return company
    company = Company(legal_name=label, normalized_name=key)
    db.add(company)
    db.flush()
    return company

def issue_token(user: User):
    exp = datetime.now(timezone.utc) + timedelta(minutes=TOKEN_MINUTES)
    return {"access_token": jwt.encode({"sub": str(user.id), "role": user.role, "ver": user.token_version, "exp": exp}, SECRET_KEY, algorithm=ALGORITHM), "token_type": "bearer", "expires_in": TOKEN_MINUTES * 60, "user": user_out(user)}

class UserInviteIn(BaseModel):
    email: EmailStr
    role: Literal["client", "auditor"] = "client"
    organization: str = Field(default="", max_length=200)

    @field_validator("email")
    @classmethod
    def gmail_only(cls, value):
        if not str(value).lower().endswith("@gmail.com"):
            raise ValueError("Only Gmail addresses can be enabled for TRACE")
        return value


def _google_state_pack(state: str, nonce: str, frontend_origin: str) -> str:
    payload = json.dumps({"state": state, "nonce": nonce, "frontend_origin": frontend_origin, "issued_at": int(time.time())}, separators=(",", ":")).encode()
    encoded = base64.urlsafe_b64encode(payload).decode().rstrip("=")
    signature = hmac.new(SECRET_KEY.encode(), encoded.encode(), hashlib.sha256).hexdigest()
    return f"{encoded}.{signature}"


def _google_state_unpack(value: str, expected_state: str) -> dict[str, Any] | None:
    try:
        encoded, signature = value.split(".", 1)
        expected_signature = hmac.new(SECRET_KEY.encode(), encoded.encode(), hashlib.sha256).hexdigest()
        if not hmac.compare_digest(signature, expected_signature):
            return None
        payload = json.loads(base64.urlsafe_b64decode(encoded + "=" * (-len(encoded) % 4)))
        if payload.get("state") != expected_state or time.time() - payload.get("issued_at", 0) > 600:
            return None
        return payload
    except (ValueError, TypeError, json.JSONDecodeError):
        return None


def _google_popup_response(frontend_origin: str, *, result: dict[str, Any] | None = None, error: str | None = None, status_code: int = 200) -> HTMLResponse:
    message = {"type": "trace-google-auth", "result": result, "error": error}
    serialized = json.dumps(message, separators=(",", ":")).replace("<", "\\u003c")
    html = f"""<!doctype html><meta charset="utf-8"><title>TRACE Google sign-in</title>
<p id="status">Returning to TRACE…</p><script>
const message = {serialized};
if (window.opener) {{ window.opener.postMessage(message, {json.dumps(frontend_origin)}); window.close(); }}
else {{ document.getElementById('status').textContent = message.error || 'Sign-in complete. You may close this window.'; }}
</script>"""
    response = HTMLResponse(html, status_code=status_code, headers={"Cache-Control": "no-store", "Referrer-Policy": "no-referrer", "X-Content-Type-Options": "nosniff"})
    response.delete_cookie("trace_google_oauth_state", path="/api/v1/auth/google")
    return response

class CompanyIn(BaseModel):
    legal_name: str = Field(min_length=2, max_length=240)

class EngagementIn(BaseModel):
    name: str = Field(min_length=2, max_length=200)
    organization: str = Field(min_length=1, max_length=200)
    period: str = Field(min_length=1, max_length=100)
    scope: str = Field(min_length=1, max_length=200)
    required_records: list[str] = Field(default_factory=list)

class EngagementPatch(BaseModel):
    name: str | None = Field(default=None, max_length=200)
    period: str | None = Field(default=None, max_length=100)
    scope: str | None = Field(default=None, max_length=200)

class AssignIn(BaseModel):
    auditor_id: int
    allowed_fields: list[str] = Field(default_factory=list)

class GrantsIn(BaseModel):
    allowed_fields: list[str]

class CheckIn(BaseModel):
    procedure: str = Field(min_length=2, max_length=100)
    version: str = "1.0"

class FindingIn(BaseModel):
    title: str = Field(min_length=2, max_length=240)
    description: str = ""
    severity: Literal["low", "medium", "high", "critical"] = "medium"

class ReviewIn(BaseModel):
    finding_id: int
    decision: Literal["accepted", "rejected", "needs_more_evidence"]
    note: str = ""

class AccountActiveIn(BaseModel):
    active: bool

class EvidenceRequestIn(BaseModel):
    description: str = Field(min_length=3, max_length=2000)
    due_date: date | None = None

class EvidenceRequestUpdate(BaseModel):
    status: Literal["submitted", "accepted", "rejected", "fulfilled"]
    note: str = Field(default="", max_length=4000)
    artifact_id: str | None = Field(default=None, max_length=40)

from fastapi.security import HTTPBearer, HTTPAuthorizationCredentials
bearer = HTTPBearer()
def authenticated_user(credentials: HTTPAuthorizationCredentials = Depends(bearer), db: Session = Depends(get_db)) -> User:
    try:
        payload = jwt.decode(credentials.credentials, SECRET_KEY, algorithms=[ALGORITHM])
        user_id = int(payload.get("sub", ""))
    except (JWTError, ValueError):
        raise HTTPException(status_code=401, detail="Invalid or expired access token", headers={"WWW-Authenticate": "Bearer"})
    user = db.get(User, user_id)
    if not user or not user.active or payload.get("ver", 0) != user.token_version:
        raise HTTPException(status_code=401, detail="Account is unavailable", headers={"WWW-Authenticate": "Bearer"})
    if user.role == "client" and user.company_id:
        company = db.get(Company, user.company_id)
        if not company or not company.active:
            raise HTTPException(status_code=401, detail="Company account is unavailable", headers={"WWW-Authenticate": "Bearer"})
    return user

UserDep = Annotated[User, Depends(authenticated_user)]
DbDep = Annotated[Session, Depends(get_db)]

def require_role(user: User, *roles: str):
    if user.role not in roles:
        raise HTTPException(status_code=403, detail="This action is not available for your role")

def get_engagement(db: Session, public_id: str) -> Engagement:
    engagement = db.scalar(select(Engagement).where(Engagement.public_id == public_id))
    if not engagement: raise HTTPException(404, "Engagement not found")
    return engagement

def can_view_engagement(db: Session, user: User, engagement: Engagement):
    # Platform admins may see operational metadata but no artifact content or extracted records.
    if user.role == "admin": return
    if user.role == "client" and user.company_id and user.company_id == engagement.company_id: return
    if user.role == "auditor" and user.id == engagement.auditor_id: return
    raise HTTPException(404, "Engagement not found")

def can_view_artifact(db: Session, user: User, engagement: Engagement):
    can_view_engagement(db, user, engagement)
    if user.role == "admin": raise HTTPException(403, "Administrators cannot access evidence documents")
    if user.role == "auditor" and not db.scalar(select(AccessGrant.id).where(AccessGrant.engagement_id == engagement.id, AccessGrant.auditor_id == user.id, AccessGrant.field_name == "documents", AccessGrant.allowed.is_(True))):
        raise HTTPException(403, "The client has not granted document access")

ACCESS_FIELDS = {"organization", "period", "scope", "required_records", "documents", "transactions", "findings", "events", "report", "checks", "sampling", "compliance"}
def require_grant(db: Session, user: User, engagement: Engagement, field: str):
    can_view_engagement(db, user, engagement)
    if user.role == "admin" and field != "events":
        raise HTTPException(403, "Platform administrators cannot access client audit data")
    if user.role == "auditor" and not db.scalar(select(AccessGrant.id).where(AccessGrant.engagement_id == engagement.id, AccessGrant.auditor_id == user.id, AccessGrant.field_name == field, AccessGrant.allowed.is_(True))):
        raise HTTPException(403, f"The client has not granted {field} access")

def add_event(db: Session, engagement: Engagement, actor: User, kind: str, details: dict):
    previous = db.scalar(select(AuditEvent).where(AuditEvent.engagement_id == engagement.id).order_by(AuditEvent.id.desc()).limit(1))
    previous_hash = previous.event_hash if previous else "0" * 64
    created = now()
    material = json.dumps({"engagement": engagement.public_id, "actor": actor.id, "type": kind, "details": details, "previous": previous_hash, "at": created.isoformat()}, sort_keys=True, separators=(",", ":"))
    event = AuditEvent(engagement_id=engagement.id, actor_id=actor.id, event_type=kind, detail_json=json.dumps(details), previous_hash=previous_hash, event_hash=hashlib.sha256(material.encode()).hexdigest(), created_at=created)
    db.add(event)

def engagement_out(db: Session, engagement: Engagement, user: User):
    out = {"id": engagement.public_id, "name": engagement.name, "organization": engagement.organization, "period": engagement.period, "scope": engagement.scope, "required_records": json.loads(engagement.required_records), "status": engagement.status, "client_id": engagement.client_id, "auditor_id": engagement.auditor_id, "created_at": engagement.created_at, "updated_at": engagement.updated_at}
    if user.role == "auditor":
        grants = db.scalars(select(AccessGrant).where(AccessGrant.engagement_id == engagement.id, AccessGrant.auditor_id == user.id, AccessGrant.allowed.is_(True))).all()
        allowed = {g.field_name for g in grants}
        for field in ("organization", "period", "scope", "required_records"):
            if field not in allowed: out.pop(field, None)
    return out

@app.get("/", tags=["default"])
def root(): return {"service": "TRACE API", "docs": "/docs", "health": "/health"}

@app.get("/health", tags=["health"])
def health(): return {"status": "ok", "database": "configured", "version": app.version}

@app.get("/api/v1/auth/google/start", tags=["auth"])
def google_login_start(request: Request, frontend_origin: str = Query(..., min_length=8, max_length=300)):
    if not GOOGLE_CLIENT_ID or not GOOGLE_CLIENT_SECRET:
        raise HTTPException(503, "Google sign-in is not configured on the backend")
    if frontend_origin not in origins:
        raise HTTPException(400, "Frontend origin is not allowed")
    state, nonce = secrets.token_urlsafe(32), secrets.token_urlsafe(32)
    params = {
        "client_id": GOOGLE_CLIENT_ID,
        "redirect_uri": GOOGLE_REDIRECT_URI,
        "response_type": "code",
        "scope": "openid email profile",
        "state": state,
        "nonce": nonce,
        "prompt": "select_account",
    }
    response = RedirectResponse(f"https://accounts.google.com/o/oauth2/v2/auth?{urlencode(params)}", status_code=302)
    response.set_cookie("trace_google_oauth_state", _google_state_pack(state, nonce, frontend_origin), max_age=600, httponly=True, secure=request.url.scheme == "https", samesite="lax", path="/api/v1/auth/google")
    return response


@app.get("/api/v1/auth/google/callback", tags=["auth"], response_class=HTMLResponse)
async def google_login_callback(request: Request, db: DbDep, state: str = "", code: str | None = None, error: str | None = None):
    packed_state = request.cookies.get("trace_google_oauth_state", "")
    oauth_state = _google_state_unpack(packed_state, state)
    if not oauth_state or oauth_state.get("frontend_origin") not in origins:
        raise HTTPException(400, "Google sign-in state is invalid or expired. Start sign-in again.")
    frontend_origin = oauth_state["frontend_origin"]
    if error:
        return _google_popup_response(frontend_origin, error="Google sign-in was cancelled or denied.")
    if not code:
        return _google_popup_response(frontend_origin, error="Google did not return an authorization code.")
    if not GOOGLE_CLIENT_ID or not GOOGLE_CLIENT_SECRET:
        return _google_popup_response(frontend_origin, error="Google sign-in is not configured on the backend.", status_code=503)
    try:
        async with httpx.AsyncClient(timeout=15) as client:
            token_response = await client.post("https://oauth2.googleapis.com/token", data={"code": code, "client_id": GOOGLE_CLIENT_ID, "client_secret": GOOGLE_CLIENT_SECRET, "redirect_uri": GOOGLE_REDIRECT_URI, "grant_type": "authorization_code"})
        token_response.raise_for_status()
        google_tokens = token_response.json()
        raw_id_token = google_tokens.get("id_token")
        if not raw_id_token:
            raise ValueError("Google did not return an identity token")
        from google.auth.transport.requests import Request as GoogleRequest
        from google.oauth2 import id_token as google_id_token
        claims = await run_in_threadpool(google_id_token.verify_oauth2_token, raw_id_token, GoogleRequest(), GOOGLE_CLIENT_ID)
    except Exception:
        return _google_popup_response(frontend_origin, error="Google could not verify this sign-in. Try again.")
    if not hmac.compare_digest(str(claims.get("nonce", "")), str(oauth_state["nonce"])):
        return _google_popup_response(frontend_origin, error="Google sign-in verification failed. Start sign-in again.")
    email = str(claims.get("email", "")).strip().lower()
    if claims.get("email_verified") not in (True, "true") or not email or not claims.get("sub"):
        return _google_popup_response(frontend_origin, error="Google must provide a verified email address.")
    if not email.endswith("@gmail.com"):
        return _google_popup_response(frontend_origin, error="TRACE sign-in currently supports Gmail addresses only.")
    user = db.scalar(select(User).where(User.google_subject == str(claims["sub"])))
    if user is None:
        user = db.scalar(select(User).where(User.email == email))
        if user is None:
            if email != PLATFORM_ADMIN_EMAIL:
                return _google_popup_response(frontend_origin, error="This Gmail is not enabled for TRACE. Ask the platform administrator to add it first.")
            user = User(email=email, name=str(claims.get("name") or "TRACE Administrator")[:200], role="admin", organization="TRACE")
            db.add(user)
            db.flush()
        if user.google_subject and user.google_subject != str(claims["sub"]):
            return _google_popup_response(frontend_origin, error="This TRACE account is already linked to another Google account.")
        user.google_subject = str(claims["sub"])
    if email != user.email.casefold():
        return _google_popup_response(frontend_origin, error="The Google account email does not match this TRACE account.")
    if email == PLATFORM_ADMIN_EMAIL:
        user.role = "admin"
        user.organization = "TRACE"
        user.active = True
    elif user.role == "admin":
        return _google_popup_response(frontend_origin, error="Administrator access is restricted to the designated Google account.")
    user.name = str(claims.get("name") or user.name)[:200]
    if not user.active:
        return _google_popup_response(frontend_origin, error="This TRACE account is disabled.")
    source_ip = request.client.host if request.client else ""
    user.last_login_at = now()
    db.add(SecurityEvent(actor_id=user.id, email=user.email, event_type="auth.google_login_succeeded", source_ip=source_ip))
    db.commit()
    return _google_popup_response(frontend_origin, result=issue_token(user))

@app.get("/api/v1/auth/me", tags=["auth"])
def me(user: UserDep): return user_out(user)

@app.post("/api/v1/admin/users", status_code=201, tags=["admin"])
def create_user(body: UserInviteIn, admin: UserDep, db: DbDep):
    require_role(admin, "admin")
    email = body.email.lower().strip()
    if not email.endswith("@gmail.com"):
        raise HTTPException(422, "Only Gmail addresses can be enabled for TRACE")
    if email == PLATFORM_ADMIN_EMAIL:
        raise HTTPException(409, "The platform administrator account is provisioned through Google sign-in.")
    company = get_or_create_company(db, body.organization) if body.role == "client" else None
    user = db.scalar(select(User).where(User.email == email))
    if user:
        user.role = body.role
        user.organization = body.organization.strip()
        user.company_id = company.id if company else None
        user.active = True
        user.token_version += 1
        action = "user.enabled"
    else:
        user = User(email=email, name=email.split("@", 1)[0], role=body.role, organization=body.organization.strip(), company_id=company.id if company else None)
        db.add(user)
        db.flush()
        action = "user.created"
    db.add(AdminAuditEvent(actor_id=admin.id, action=action, target_type="user", target_id=str(user.id))); db.commit(); db.refresh(user)
    return user_out(user)

@app.get("/api/v1/admin/users", tags=["admin"])
def list_users(admin: UserDep, db: DbDep, role: str | None = None):
    require_role(admin, "admin")
    query = select(User).order_by(User.created_at.desc())
    if role: query = query.where(User.role == role)
    return [user_out(u) for u in db.scalars(query).all()]

@app.patch("/api/v1/admin/users/{user_id}/active", tags=["admin"])
def set_user_active(user_id: int, body: AccountActiveIn, admin: UserDep, db: DbDep):
    require_role(admin, "admin")
    target = db.get(User, user_id)
    if not target: raise HTTPException(404, "User not found")
    if target.id == admin.id and not body.active: raise HTTPException(409, "Administrators cannot disable their own account")
    target.active = body.active; target.token_version += 1
    db.add(AdminAuditEvent(actor_id=admin.id, action="user.activation_changed", target_type="user", target_id=str(target.id)))
    db.add(SecurityEvent(actor_id=admin.id, email=target.email, event_type="auth.account_activation_changed")); db.commit()
    return user_out(target)

@app.post("/api/v1/admin/companies", status_code=201, tags=["admin"])
def create_company(body: CompanyIn, admin: UserDep, db: DbDep):
    require_role(admin, "admin")
    company = get_or_create_company(db, body.legal_name)
    db.add(AdminAuditEvent(actor_id=admin.id, action="company.created", target_type="company", target_id=str(company.id)))
    db.commit()
    return {"id": company.id, "legal_name": company.legal_name, "active": company.active}

@app.get("/api/v1/admin/companies", tags=["admin"])
def list_companies(admin: UserDep, db: DbDep):
    require_role(admin, "admin")
    companies = db.scalars(select(Company).order_by(Company.legal_name)).all()
    return [{"id": c.id, "legal_name": c.legal_name, "active": c.active} for c in companies]

@app.get("/api/v1/admin/audit-logs", tags=["admin"])
def admin_audit_logs(admin: UserDep, db: DbDep):
    require_role(admin, "admin")
    return [{"id": event.id, "action": event.action, "target_type": event.target_type, "target_id": event.target_id, "created_at": event.created_at, "actor_id": event.actor_id} for event in db.scalars(select(AdminAuditEvent).order_by(AdminAuditEvent.id.desc()).limit(500)).all()]

@app.get("/api/v1/auditors", tags=["engagements"])
def list_available_auditors(user: UserDep, db: DbDep):
    require_role(user, "client")
    return [{"id": u.id, "name": u.name, "organization": u.organization} for u in db.scalars(select(User).where(User.role == "auditor", User.active.is_(True)).order_by(User.name)).all()]

@app.post("/api/v1/engagements/", status_code=201, tags=["engagements"])
def create_engagement(body: EngagementIn, user: UserDep, db: DbDep):
    require_role(user, "client")
    if not user.company_id:
        raise HTTPException(409, "Client account is not linked to a company")
    company = db.get(Company, user.company_id)
    if not company or not company.active:
        raise HTTPException(403, "Company account is inactive")
    e = Engagement(public_id=f"TR-{datetime.now().year}-{secrets.token_hex(3).upper()}", name=body.name.strip(), organization=company.legal_name, period=body.period.strip(), scope=body.scope.strip(), required_records=json.dumps(body.required_records), client_id=user.id, company_id=company.id)
    db.add(e); db.flush(); add_event(db, e, user, "engagement.created", {"scope": e.scope}); db.commit(); db.refresh(e)
    return engagement_out(db, e, user)

@app.get("/api/v1/engagements/", tags=["engagements"])
def list_engagements(user: UserDep, db: DbDep):
    query = select(Engagement).order_by(Engagement.updated_at.desc())
    if user.role == "client": query = query.where(Engagement.company_id == user.company_id)
    elif user.role == "auditor": query = query.where(Engagement.auditor_id == user.id)
    # Admin listing is metadata only.
    return [engagement_out(db, e, user) for e in db.scalars(query).all()]

@app.get("/api/v1/engagements/{engagement_id}", tags=["engagements"])
def read_engagement(engagement_id: str, user: UserDep, db: DbDep):
    e = get_engagement(db, engagement_id); can_view_engagement(db, user, e); return engagement_out(db, e, user)

@app.patch("/api/v1/engagements/{engagement_id}", tags=["engagements"])
def update_engagement(engagement_id: str, body: EngagementPatch, user: UserDep, db: DbDep):
    e = get_engagement(db, engagement_id)
    if user.role != "client" or user.id != e.client_id: raise HTTPException(403, "Only the engagement client can update it")
    changes = body.model_dump(exclude_unset=True)
    for k, v in changes.items(): setattr(e, k, v)
    add_event(db, e, user, "engagement.updated", changes); db.commit(); db.refresh(e)
    return engagement_out(db, e, user)

@app.post("/api/v1/engagements/{engagement_id}/assign", tags=["engagements"])
def assign_auditor(engagement_id: str, body: AssignIn, user: UserDep, db: DbDep):
    e = get_engagement(db, engagement_id)
    if user.id != e.client_id or user.role != "client": raise HTTPException(403, "Only the client can choose the auditor")
    auditor = db.get(User, body.auditor_id)
    if not auditor or auditor.role != "auditor" or not auditor.active: raise HTTPException(404, "Active auditor not found")
    if set(body.allowed_fields) - ACCESS_FIELDS: raise HTTPException(422, f"Unknown access fields: {sorted(set(body.allowed_fields) - ACCESS_FIELDS)}")
    e.auditor_id = auditor.id; e.status = "Assigned"
    db.query(AccessGrant).filter_by(engagement_id=e.id, auditor_id=auditor.id).delete()
    for field in set(body.allowed_fields): db.add(AccessGrant(engagement_id=e.id, auditor_id=auditor.id, field_name=field, allowed=True))
    add_event(db, e, user, "engagement.auditor_assigned", {"auditor_id": auditor.id, "allowed_fields": sorted(set(body.allowed_fields))}); db.commit()
    return {"engagement": engagement_out(db, e, user), "auditor": {"id": auditor.id, "name": auditor.name, "email": auditor.email}, "allowed_fields": sorted(set(body.allowed_fields))}

@app.put("/api/v1/engagements/{engagement_id}/access", tags=["engagements"])
def update_access(engagement_id: str, body: GrantsIn, user: UserDep, db: DbDep):
    e = get_engagement(db, engagement_id)
    if user.id != e.client_id or user.role != "client": raise HTTPException(403, "Only the client can change auditor access")
    if not e.auditor_id: raise HTTPException(409, "Assign an auditor first")
    if set(body.allowed_fields) - ACCESS_FIELDS: raise HTTPException(422, f"Unknown access fields: {sorted(set(body.allowed_fields) - ACCESS_FIELDS)}")
    db.query(AccessGrant).filter_by(engagement_id=e.id, auditor_id=e.auditor_id).delete()
    for field in set(body.allowed_fields): db.add(AccessGrant(engagement_id=e.id, auditor_id=e.auditor_id, field_name=field, allowed=True))
    add_event(db, e, user, "engagement.access_updated", {"allowed_fields": sorted(set(body.allowed_fields))}); db.commit()
    return {"allowed_fields": sorted(set(body.allowed_fields))}

def _profile_bytes(filename: str, raw: bytes):
    result = {"format": Path(filename).suffix.lower().lstrip("."), "bytes": len(raw), "sha256": hashlib.sha256(raw).hexdigest(), "rows": None, "columns": [], "missing_values": {}, "numeric_summary": {}}
    if Path(filename).suffix.lower() not in SUPPORTED_EXTENSIONS:
        result["status"] = "stored_unparsed"
        return result
    try:
        df, source = load_dataframe(filename, raw)
        profile = profile_dataframe(df, source)
        result.update({"status": "profiled", "rows": profile["row_count"], "columns": profile["columns"], "missing_values": profile["missing_values"], "numeric_summary": profile["numeric_summary"], "profile": profile})
    except DatasetError as exc:
        result["parse_error"] = str(exc); return result
    return result

@app.post("/api/v1/uploads/{engagement_id}/upload", status_code=201, tags=["uploads"])
async def upload_file(engagement_id: str, user: UserDep, db: DbDep, file: UploadFile = File(...)):
    e = get_engagement(db, engagement_id); can_view_artifact(db, user, e)
    if user.role not in {"client", "auditor"}: raise HTTPException(403, "Role cannot upload evidence")
    safe_name = Path(file.filename or "upload").name[:255]
    extension = Path(safe_name).suffix.lower()
    allowed = SUPPORTED_EXTENSIONS | {".pdf", ".docx", ".doc"}
    if extension not in allowed: raise HTTPException(415, f"Unsupported evidence format. Use: {', '.join(sorted(allowed))}")
    raw = await file.read(MAX_UPLOAD_BYTES + 1)
    if len(raw) > MAX_UPLOAD_BYTES: raise HTTPException(413, f"File exceeds {MAX_UPLOAD_BYTES} bytes")
    if not raw: raise HTTPException(400, "Uploaded file is empty")
    if extension in SUPPORTED_EXTENSIONS:
        try: load_dataframe(safe_name, raw)
        except DatasetError as exc: raise HTTPException(422, str(exc))
    elif extension == ".pdf" and not raw.startswith(b"%PDF-"):
        raise HTTPException(422, "File extension does not match a PDF document")
    elif extension in {".docx", ".xlsx", ".xlsm"} and not raw.startswith(b"PK\x03\x04"):
        raise HTTPException(422, "File extension does not match a valid Office document")
    digest = hashlib.sha256(raw).hexdigest(); artifact_id = str(uuid4())
    dest = UPLOAD_DIR / f"{artifact_id}{Path(safe_name).suffix.lower()}"
    dest.write_bytes(raw)
    try: dest.chmod(0o600)
    except OSError: pass
    artifact = Artifact(public_id=artifact_id, engagement_id=e.id, uploader_id=user.id, filename=safe_name, content_type=file.content_type or "application/octet-stream", storage_path=str(dest), sha256=digest, size=len(raw), validation_json=json.dumps({"status": "uploaded", "profile": _profile_bytes(safe_name, raw)}))
    db.add(artifact); add_event(db, e, user, "evidence.uploaded", {"artifact_id": artifact_id, "filename": safe_name, "sha256": digest, "size": len(raw)}); db.commit()
    return artifact_out(artifact)

def artifact_out(a: Artifact):
    return {"id": a.public_id, "filename": a.filename, "content_type": a.content_type, "sha256": a.sha256, "size": a.size, "validation": json.loads(a.validation_json), "created_at": a.created_at}

@app.get("/api/v1/uploads/{engagement_id}/uploads", tags=["uploads"])
def list_uploads(engagement_id: str, user: UserDep, db: DbDep):
    e = get_engagement(db, engagement_id); can_view_artifact(db, user, e)
    return [artifact_out(a) for a in db.scalars(select(Artifact).where(Artifact.engagement_id == e.id).order_by(Artifact.created_at.desc())).all()]

@app.get("/api/v1/uploads/{engagement_id}/uploads/{artifact_id}/download", tags=["uploads"])
def download_upload(engagement_id: str, artifact_id: str, user: UserDep, db: DbDep):
    e = get_engagement(db, engagement_id); can_view_artifact(db, user, e)
    a = db.scalar(select(Artifact).where(Artifact.engagement_id == e.id, Artifact.public_id == artifact_id))
    if not a: raise HTTPException(404, "Artifact not found")
    return Response(Path(a.storage_path).read_bytes(), media_type=a.content_type, headers={"Content-Disposition": f'attachment; filename="{a.filename}"', "X-Content-SHA256": a.sha256})

@app.post("/api/v1/uploads/{engagement_id}/validate/{artifact_id}", tags=["uploads"])
def validate_upload(engagement_id: str, artifact_id: str, user: UserDep, db: DbDep):
    e = get_engagement(db, engagement_id); can_view_artifact(db, user, e)
    a = db.scalar(select(Artifact).where(Artifact.engagement_id == e.id, Artifact.public_id == artifact_id))
    if not a: raise HTTPException(404, "Artifact not found")
    raw = Path(a.storage_path).read_bytes(); profile = _profile_bytes(a.filename, raw)
    if hashlib.sha256(raw).hexdigest() != a.sha256: raise HTTPException(409, "Stored evidence integrity check failed")
    payload = {"status": "validated", "sha256_verified": True, "profile": profile, "validated_at": now().isoformat()}
    a.validation_json = json.dumps(payload); add_event(db, e, user, "evidence.validated", {"artifact_id": a.public_id, "sha256": a.sha256, "rows": profile.get("rows"), "columns": profile.get("columns")}); db.commit()
    return {"artifact_id": a.public_id, **payload}

def _xlsx_rows(raw: bytes):
    from openpyxl import load_workbook
    workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    try:
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        first = next(iterator, None)
        if first is None: return [], []
        columns = [str(v).strip() if v is not None else "" for v in first]
        columns = [col or f"column_{i+1}" for i, col in enumerate(columns)]
        rows = []
        for values in iterator:
            row = {}
            for index, column in enumerate(columns):
                value = values[index] if index < len(values) else None
                row[column] = value.isoformat() if hasattr(value, "isoformat") else "" if value is None else str(value)
            if any(value.strip() for value in row.values()): rows.append(row)
        return columns, rows
    finally:
        workbook.close()

def csv_dataset_for_engagement(db: Session, e: Engagement):
    artifacts = db.scalars(select(Artifact).where(Artifact.engagement_id == e.id).order_by(Artifact.created_at.desc())).all()
    for a in artifacts:
        if Path(a.filename).suffix.lower() in SUPPORTED_EXTENSIONS:
            try:
                raw = Path(a.storage_path).read_bytes()
                df, _ = load_dataframe(a.filename, raw)
                return a, list(df.columns), dataframe_records(df)
            except (OSError, DatasetError): continue
    return None, [], []

@app.post("/api/v1/audit/{engagement_id}/snapshots", status_code=201, tags=["audit"])
def create_snapshot(engagement_id: str, user: UserDep, db: DbDep):
    e = get_engagement(db, engagement_id); can_view_artifact(db, user, e)
    a, columns, rows = csv_dataset_for_engagement(db, e)
    if not a: raise HTTPException(409, "Upload a supported tabular dataset before creating a snapshot")
    profile = _profile_bytes(a.filename, Path(a.storage_path).read_bytes()).get("profile", {})
    snapshot_id = str(uuid4())
    parent = db.scalar(select(AuditSnapshot).where(AuditSnapshot.engagement_id == e.id).order_by(AuditSnapshot.id.desc()).limit(1))
    snapshot = AuditSnapshot(engagement_id=e.id, actor_id=user.id, artifact_id=a.public_id, dataset_hash=a.sha256, row_count=len(rows), columns_json=json.dumps(columns), profile_json=json.dumps(profile), parent_snapshot_id=parent.id if parent else None)
    db.add(snapshot)
    add_event(db, e, user, "audit.snapshot_created", {"artifact_id": a.public_id, "sha256": a.sha256, "parent_snapshot_id": snapshot.parent_snapshot_id, "rows": profile.get("rows"), "columns": columns}); db.commit()
    return {"id": snapshot.id, "artifact_id": a.public_id, "dataset_hash": a.sha256, "parent_snapshot_id": snapshot.parent_snapshot_id, "row_count": len(rows), "columns": columns, "profile": profile, "created_at": snapshot.created_at}

@app.post("/api/v1/audit/{engagement_id}/checks", status_code=201, tags=["audit"])
def run_check(engagement_id: str, body: CheckIn, user: UserDep, db: DbDep):
    require_role(user, "auditor", "client")
    e = get_engagement(db, engagement_id); require_grant(db, user, e, "checks"); can_view_artifact(db, user, e)
    a, columns, rows = csv_dataset_for_engagement(db, e)
    if not a: raise HTTPException(409, "Upload a CSV or TSV dataset before running checks")
    key = body.procedure.casefold()
    check_id = key if key in CHECKS else "duplicate_invoice" if "duplicate" in key else "negative_amount" if any(term in key for term in ("financial", "amount", "numeric")) else "missing_values" if any(term in key for term in ("missing", "completeness")) else "out_of_period" if any(term in key for term in ("cutoff", "cut-off", "period")) else "duplicate_rows"
    df, _ = load_dataframe(a.filename, Path(a.storage_path).read_bytes())
    check = execute_check(check_id, df, source_artifact_id=a.public_id, params={"columns": columns})
    findings = [check]
    result = {"procedure": body.procedure, "version": body.version, "dataset_artifact_id": a.public_id, "dataset_sha256": a.sha256, "rows_examined": len(rows), "columns_examined": len(columns), "results": findings, "risk_indicators": risk_indicators(df), "note": "Deterministic data checks; auditor review is required and these results are not an audit opinion."}
    run = ProcedureRun(engagement_id=e.id, actor_id=user.id, procedure=body.procedure, version=body.version, result_json=json.dumps(result)); db.add(run)
    add_event(db, e, user, "audit.check_executed", {"procedure": body.procedure, "version": body.version, "dataset_sha256": a.sha256, "finding_count": len(findings)}); db.commit(); db.refresh(run)
    return {"id": run.id, **result, "created_at": run.created_at}

@app.get("/api/v1/audit/checks/catalog", tags=["audit"])
def check_catalog():
    return [{"id": check_id, **definition} for check_id, definition in CHECKS.items()]

@app.get("/api/v1/audit/{engagement_id}/checks", tags=["audit"])
def list_checks(engagement_id: str, user: UserDep, db: DbDep):
    e = get_engagement(db, engagement_id); require_grant(db, user, e, "checks")
    return [{"id": x.id, **json.loads(x.result_json), "created_at": x.created_at} for x in db.scalars(select(ProcedureRun).where(ProcedureRun.engagement_id == e.id).order_by(ProcedureRun.id.desc())).all()]

@app.post("/api/v1/audit/{engagement_id}/findings", status_code=201, tags=["audit"])
def create_finding(engagement_id: str, body: FindingIn, user: UserDep, db: DbDep):
    require_role(user, "auditor", "client")
    e = get_engagement(db, engagement_id); require_grant(db, user, e, "findings")
    f = Finding(engagement_id=e.id, actor_id=user.id, title=body.title, description=body.description, severity=body.severity)
    db.add(f); db.flush(); add_event(db, e, user, "finding.created", {"finding_id": f.id, "severity": f.severity}); db.commit(); db.refresh(f)
    return finding_out(f)

def finding_out(f): return {"id": f.id, "title": f.title, "description": f.description, "severity": f.severity, "status": f.status, "created_at": f.created_at}

@app.get("/api/v1/audit/{engagement_id}/findings", tags=["audit"])
def list_findings(engagement_id: str, user: UserDep, db: DbDep):
    e = get_engagement(db, engagement_id); require_grant(db, user, e, "findings")
    return [finding_out(f) for f in db.scalars(select(Finding).where(Finding.engagement_id == e.id).order_by(Finding.id.desc())).all()]

@app.post("/api/v1/audit/{engagement_id}/evidence-requests", status_code=201, tags=["audit"])
def create_evidence_request(engagement_id: str, body: EvidenceRequestIn, user: UserDep, db: DbDep):
    require_role(user, "auditor")
    e = get_engagement(db, engagement_id)
    if e.auditor_id != user.id: raise HTTPException(404, "Engagement not found")
    require_grant(db, user, e, "findings")
    request = EvidenceRequest(engagement_id=e.id, actor_id=user.id, description=body.description.strip(), due_date=body.due_date)
    db.add(request); db.flush(); add_event(db, e, user, "evidence.requested", {"request_id": request.id})
    db.commit(); db.refresh(request)
    return {"id": request.id, "description": request.description, "status": request.status, "due_date": request.due_date, "created_at": request.created_at}

@app.get("/api/v1/audit/{engagement_id}/evidence-requests", tags=["audit"])
def list_evidence_requests(engagement_id: str, user: UserDep, db: DbDep):
    e = get_engagement(db, engagement_id); can_view_engagement(db, user, e)
    if user.role == "auditor": require_grant(db, user, e, "findings")
    document_granted = user.role != "auditor" or bool(db.scalar(select(AccessGrant.id).where(AccessGrant.engagement_id == e.id, AccessGrant.auditor_id == user.id, AccessGrant.field_name == "documents", AccessGrant.allowed.is_(True))))
    return [{"id": x.id, "description": x.description, "status": x.status, "due_date": x.due_date, "artifact_id": x.artifact_public_id if document_granted else None, "response_note": x.response_note, "created_at": x.created_at, "updated_at": x.updated_at} for x in db.scalars(select(EvidenceRequest).where(EvidenceRequest.engagement_id == e.id).order_by(EvidenceRequest.id.desc())).all()]

@app.patch("/api/v1/audit/{engagement_id}/evidence-requests/{request_id}", tags=["audit"])
def update_evidence_request(engagement_id: str, request_id: int, body: EvidenceRequestUpdate, user: UserDep, db: DbDep):
    e = get_engagement(db, engagement_id); can_view_engagement(db, user, e)
    request = db.scalar(select(EvidenceRequest).where(EvidenceRequest.id == request_id, EvidenceRequest.engagement_id == e.id))
    if not request: raise HTTPException(404, "Evidence request not found")
    if user.role == "client":
        if body.status != "submitted" or request.status not in {"open", "rejected"}: raise HTTPException(409, "Client responses must submit an open or rejected request")
        if body.artifact_id:
            artifact = db.scalar(select(Artifact).where(Artifact.engagement_id == e.id, Artifact.public_id == body.artifact_id))
            if not artifact: raise HTTPException(404, "Response evidence not found")
            request.artifact_public_id = artifact.public_id
        request.status = "submitted"; request.response_note = body.note.strip()
        event_type = "evidence.submitted"
    elif user.role == "auditor":
        if e.auditor_id != user.id: raise HTTPException(404, "Engagement not found")
        require_grant(db, user, e, "findings")
        if body.status not in {"accepted", "rejected"} or request.status != "submitted": raise HTTPException(409, "Auditors can decide only on submitted evidence")
        request.status = body.status; request.response_note = body.note.strip(); event_type = "evidence.reviewed"
    else:
        raise HTTPException(403, "Role cannot update evidence requests")
    request.updated_at = now()
    add_event(db, e, user, event_type, {"request_id": request.id, "status": request.status, "artifact_id": request.artifact_public_id})
    db.commit()
    return {"id": request.id, "status": request.status, "artifact_id": request.artifact_public_id, "response_note": request.response_note, "updated_at": request.updated_at}

@app.post("/api/v1/audit/{engagement_id}/review-decisions", status_code=201, tags=["audit"])
def review_finding(engagement_id: str, body: ReviewIn, user: UserDep, db: DbDep):
    require_role(user, "auditor")
    e = get_engagement(db, engagement_id)
    if e.auditor_id != user.id: raise HTTPException(403, "Assigned auditor only")
    require_grant(db, user, e, "findings")
    f = db.get(Finding, body.finding_id)
    if not f or f.engagement_id != e.id: raise HTTPException(404, "Finding not found")
    f.status = body.decision; add_event(db, e, user, "finding.reviewed", {"finding_id": f.id, "decision": body.decision, "note": body.note}); db.commit()
    return {"finding_id": f.id, "decision": f.status, "note": body.note}

@app.get("/api/v1/audit/{engagement_id}/audit-events", tags=["audit"])
def list_events(engagement_id: str, user: UserDep, db: DbDep):
    e = get_engagement(db, engagement_id); require_grant(db, user, e, "events")
    query = select(AuditEvent).where(AuditEvent.engagement_id == e.id)
    if user.role == "admin": query = query.where(AuditEvent.actor_id == user.id)
    rows = db.scalars(query.order_by(AuditEvent.id.desc())).all()
    # Admin receives only abstract event metadata; event details can include sensitive field names.
    return [{"id": x.id, "type": x.event_type, "actor_id": x.actor_id, "created_at": x.created_at, "previous_hash": x.previous_hash, "event_hash": x.event_hash, "details": {} if user.role == "admin" else json.loads(x.detail_json)} for x in rows]

@app.get("/api/v1/audit/{engagement_id}/audit-integrity", tags=["audit"])
def verify_integrity(engagement_id: str, user: UserDep, db: DbDep):
    e = get_engagement(db, engagement_id); require_grant(db, user, e, "events")
    rows = db.scalars(select(AuditEvent).where(AuditEvent.engagement_id == e.id).order_by(AuditEvent.id.asc())).all()
    expected_prev = "0" * 64; valid = True
    for event in rows:
        detail = json.loads(event.detail_json)
        event_time = event.created_at.replace(tzinfo=timezone.utc) if event.created_at.tzinfo is None else event.created_at
        material = json.dumps({"engagement": e.public_id, "actor": event.actor_id, "type": event.event_type, "details": detail, "previous": expected_prev, "at": event_time.isoformat()}, sort_keys=True, separators=(",", ":"))
        expected_hash = hashlib.sha256(material.encode()).hexdigest()
        if event.previous_hash != expected_prev or event.event_hash != expected_hash: valid = False; break
        expected_prev = event.event_hash
    return {"engagement_id": e.public_id, "valid": valid, "events_checked": len(rows), "head_hash": expected_prev, "algorithm": "SHA-256 chained events", "note": "Tamper-evident application log; not a blockchain or independently immutable ledger."}

@app.get("/api/v1/audit/{engagement_id}/report", tags=["audit"])
def generate_report(engagement_id: str, user: UserDep, db: DbDep):
    e = get_engagement(db, engagement_id); require_grant(db, user, e, "report")
    runs = db.scalars(select(ProcedureRun).where(ProcedureRun.engagement_id == e.id).order_by(ProcedureRun.id.asc())).all()
    findings = db.scalars(select(Finding).where(Finding.engagement_id == e.id).order_by(Finding.id.asc())).all()
    artifacts = db.scalars(select(Artifact).where(Artifact.engagement_id == e.id).order_by(Artifact.created_at.asc())).all()
    profile = json.loads(artifacts[-1].validation_json).get("profile", {}) if artifacts else {}
    revenue_analysis = analyze_revenue_datasets(artifacts)
    result = {"engagement": engagement_out(db, e, user), "generated_at": now().isoformat(), "dataset_versions": [{"artifact_id": a.public_id, "filename": a.filename, "sha256": a.sha256, "uploaded_at": a.created_at} for a in artifacts], "latest_dataset_profile": profile, "revenue_analysis": revenue_analysis, "procedure_runs": [{"id": r.id, **json.loads(r.result_json), "created_at": r.created_at} for r in runs], "findings": [finding_out(f) for f in findings], "human_review_required": True, "disclaimer": "This generated workpaper summarizes uploaded data and deterministic checks. It is not an audit opinion or assurance conclusion."}
    latest_snapshot = db.scalar(select(AuditSnapshot).where(AuditSnapshot.engagement_id == e.id).order_by(AuditSnapshot.id.desc()).limit(1))
    latest_artifact_hash = artifacts[-1].sha256 if artifacts else None
    snapshot_id = latest_snapshot.id if latest_snapshot and latest_snapshot.dataset_hash == latest_artifact_hash else None
    prior_versions = db.scalars(select(ReportVersion.version).where(ReportVersion.engagement_id == e.id)).all()
    report_version = max(prior_versions, default=0) + 1
    result["report_version"] = report_version
    result["snapshot_id"] = snapshot_id
    serialized = json.dumps(result, default=str, sort_keys=True, separators=(",", ":"))
    version = ReportVersion(engagement_id=e.id, actor_id=user.id, version=report_version, snapshot_id=snapshot_id, report_hash=hashlib.sha256(serialized.encode()).hexdigest(), report_json=serialized)
    db.add(version)
    add_event(db, e, user, "report.generated", {"version": report_version, "snapshot_id": snapshot_id, "report_hash": version.report_hash, "procedure_runs": len(runs), "findings": len(findings), "dataset_versions": len(artifacts)}); db.commit()
    return result

@app.get("/api/v1/audit/{engagement_id}/report/versions", tags=["audit"])
def report_versions(engagement_id: str, user: UserDep, db: DbDep):
    e = get_engagement(db, engagement_id); require_grant(db, user, e, "report")
    versions = db.scalars(select(ReportVersion).where(ReportVersion.engagement_id == e.id).order_by(ReportVersion.version.desc())).all()
    return [{"version": r.version, "snapshot_id": r.snapshot_id, "report_hash": r.report_hash, "created_at": r.created_at} for r in versions]

@app.post("/api/v1/audit/{engagement_id}/scan", tags=["audit"])
def continue_scan(engagement_id: str, user: UserDep, db: DbDep, key_column: str | None = None):
    require_role(user, "client", "auditor")
    e = get_engagement(db, engagement_id); can_view_artifact(db, user, e)
    artifacts = db.scalars(select(Artifact).where(Artifact.engagement_id == e.id).order_by(Artifact.created_at.desc()).limit(2)).all()
    if len(artifacts) < 2: raise HTTPException(409, "Upload a revised dataset version to compare against the previous version")
    latest, baseline = artifacts[0], artifacts[1]
    try:
        def dataset_rows(artifact):
            raw = Path(artifact.storage_path).read_bytes()
            dataframe, _ = load_dataframe(artifact.filename, raw)
            return list(dataframe.columns), dataframe_records(dataframe)
        old_columns, old_rows = dataset_rows(baseline); new_columns, new_rows = dataset_rows(latest)
    except (DatasetError, OSError): raise HTTPException(422, "Both evidence versions must be readable supported tabular datasets")
    if old_columns != new_columns: raise HTTPException(422, "Dataset columns differ; align the schema before scanning")
    if key_column and key_column not in old_columns: raise HTTPException(422, "The selected key column is not present in both datasets")
    def row_key(row, index):
        values = [row.get(key_column, "")] if key_column else [row.get(c, "") for c in old_columns]
        return hashlib.sha256(json.dumps(values, sort_keys=True, ensure_ascii=False).encode()).hexdigest(), index
    old_map = {row_key(row, i)[0]: row for i, row in enumerate(old_rows)}
    new_map = {row_key(row, i)[0]: row for i, row in enumerate(new_rows)}
    old_keys, new_keys = set(old_map), set(new_map)
    # With a selected stable key, detect changed rows under that key as well as additions/removals.
    changed = []
    if key_column:
        old_map = {str(r.get(key_column, "")): r for r in old_rows}; new_map = {str(r.get(key_column, "")): r for r in new_rows}
        old_keys, new_keys = set(old_map), set(new_map)
        changed = [k for k in old_keys & new_keys if old_map[k] != new_map[k]]
    try:
        previous_df, _ = load_dataframe(baseline.filename, Path(baseline.storage_path).read_bytes())
        current_df, current_source = load_dataframe(latest.filename, Path(latest.storage_path).read_bytes())
        previous_checks = run_checks(previous_df, source_artifact_id=baseline.public_id)
        current_checks = run_checks(current_df, source_artifact_id=latest.public_id)
    except DatasetError:
        raise HTTPException(422, "Could not reprocess both versions for exception comparison")
    def stable_exceptions(checks):
        result_set = set()
        for check in checks:
            for reference in check.get("source_record_references", []):
                if reference.get("permanent_key", True) is False: continue
                if reference.get("key") is not None:
                    result_set.add((check["check_id"], reference.get("key_column"), reference["key"]))
        return result_set
    old_exceptions, new_exceptions = stable_exceptions(previous_checks), stable_exceptions(current_checks)
    parent = db.scalar(select(AuditSnapshot).where(AuditSnapshot.engagement_id == e.id).order_by(AuditSnapshot.id.desc()).limit(1))
    current_snapshot = parent if parent and parent.dataset_hash == latest.sha256 else None
    if current_snapshot is None:
        profile = profile_dataframe(current_df, {**current_source, "artifact_id": latest.public_id, "sha256": latest.sha256})
        current_snapshot = AuditSnapshot(engagement_id=e.id, actor_id=user.id, artifact_id=latest.public_id, dataset_hash=latest.sha256, row_count=len(current_df), columns_json=json.dumps(list(current_df.columns)), profile_json=json.dumps(profile), parent_snapshot_id=parent.id if parent else None)
        db.add(current_snapshot); db.flush()
    result = {"baseline_artifact_id": baseline.public_id, "current_artifact_id": latest.public_id, "baseline_sha256": baseline.sha256, "current_sha256": latest.sha256, "baseline_snapshot_id": parent.id if parent and parent.dataset_hash == baseline.sha256 else None, "current_snapshot_id": current_snapshot.id, "added_rows": len(new_keys - old_keys), "removed_rows": len(old_keys - new_keys), "changed_rows": len(changed), "changed_keys_sample": changed[:100], "new_exception_count": len(new_exceptions - old_exceptions), "resolved_exception_count": len(old_exceptions - new_exceptions), "exception_comparison_uses_stable_business_keys": True, "key_column": key_column, "human_review_required": True, "note": "Row and exception comparisons are change signals only. Exceptions without stable business identifiers are not classified as newly detected or resolved."}
    previous_versions = db.scalars(select(ReportVersion.version).where(ReportVersion.engagement_id == e.id)).all()
    report_version = max(previous_versions, default=0) + 1
    report_payload = {"engagement_id": e.public_id, "report_version": report_version, "snapshot_id": current_snapshot.id, "continuous_scan": result, "human_review_required": True}
    report_json = json.dumps(report_payload, sort_keys=True, separators=(",", ":"))
    report_version_row = ReportVersion(engagement_id=e.id, actor_id=user.id, version=report_version, snapshot_id=current_snapshot.id, report_hash=hashlib.sha256(report_json.encode()).hexdigest(), report_json=report_json)
    db.add(report_version_row); result["report_version"] = report_version
    add_event(db, e, user, "audit.continuous_scan", result); db.commit()
    return result

@app.post("/api/v1/audit/{engagement_id}/sampling", status_code=201, tags=["audit"])
def run_sampling(engagement_id: str, user: UserDep, db: DbDep, method: str = Query("random"), size: int = Query(60, ge=1, le=10000)):
    require_role(user, "auditor", "client")
    e = get_engagement(db, engagement_id); require_grant(db, user, e, "sampling")
    a, columns, rows = csv_dataset_for_engagement(db, e)
    if not a: raise HTTPException(409, "Upload a CSV or TSV dataset first")
    # Stable selection tied to the uploaded snapshot for reproducibility.
    df, _ = load_dataframe(a.filename, Path(a.storage_path).read_bytes())
    if method not in {"random", "risk-based", "targeted"}: raise HTTPException(422, "Unsupported sampling method")
    plan = deterministic_sample(df, source_fingerprint=a.sha256, size=size, seed=f"trace-{method}-v1", strategy=method)
    selected = plan["selected_records"]
    chosen = [ref.get("source_row_number", 2) - 2 for ref in selected]
    result = {"method": method, "algorithm": plan["algorithm"], "seed": plan["seed"], "population_size": len(rows), "sample_size": len(chosen), "row_indexes": chosen, "selected_records": selected, "dataset_sha256": a.sha256, "note": "Reproducible technical selection only; sample sufficiency requires auditor judgment."}
    run = SamplingRun(engagement_id=e.id, actor_id=user.id, result_json=json.dumps(result)); db.add(run)
    add_event(db, e, user, "audit.sampling_executed", {"method": method, "population_size": len(rows), "sample_size": len(chosen), "dataset_sha256": a.sha256}); db.commit(); db.refresh(run)
    return {"id": run.id, **result, "created_at": run.created_at}

@app.get("/api/v1/audit/{engagement_id}/sampling", tags=["audit"])
def list_sampling(engagement_id: str, user: UserDep, db: DbDep):
    e = get_engagement(db, engagement_id); require_grant(db, user, e, "sampling")
    return [{"id": r.id, **json.loads(r.result_json), "created_at": r.created_at} for r in db.scalars(select(SamplingRun).where(SamplingRun.engagement_id == e.id).order_by(SamplingRun.id.desc())).all()]

@app.get("/api/v1/audit/{engagement_id}/transactions", tags=["audit"])
def list_transactions(engagement_id: str, user: UserDep, db: DbDep, limit: int = Query(100, ge=1, le=1000), offset: int = Query(0, ge=0)):
    e = get_engagement(db, engagement_id); require_grant(db, user, e, "transactions")
    a, columns, rows = csv_dataset_for_engagement(db, e)
    if not a: return {"items": [], "columns": [], "total": 0}
    return {"items": rows[offset:offset + limit], "columns": columns, "total": len(rows), "artifact_id": a.public_id}

@app.get("/api/v1/audit/{engagement_id}/compliance", tags=["audit"])
def compliance(engagement_id: str, user: UserDep, db: DbDep):
    e = get_engagement(db, engagement_id); require_grant(db, user, e, "compliance")
    return []

@app.post("/api/v1/audit/{engagement_id}/compliance", tags=["audit"])
def create_compliance(engagement_id: str, body: dict, user: UserDep, db: DbDep):
    require_role(user, "auditor")
    e = get_engagement(db, engagement_id)
    if e.auditor_id != user.id: raise HTTPException(403, "Assigned auditor only")
    require_grant(db, user, e, "compliance")
    add_event(db, e, user, "compliance.procedure_recorded", body); db.commit()
    return {"status": "recorded", "procedure": body, "human_review_required": True}

@app.get("/api/v1/audit/{engagement_id}/exports/{export_type}", tags=["audit"])
def export_workpaper(engagement_id: str, export_type: Literal["json", "csv"], user: UserDep, db: DbDep):
    e = get_engagement(db, engagement_id); require_grant(db, user, e, "report")
    granted = {g.field_name for g in db.scalars(select(AccessGrant).where(AccessGrant.engagement_id == e.id, AccessGrant.auditor_id == user.id, AccessGrant.allowed.is_(True))).all()} if user.role == "auditor" else ACCESS_FIELDS
    if export_type == "csv" and "transactions" not in granted: raise HTTPException(403, "The client has not granted transaction export access")
    a, columns, rows = csv_dataset_for_engagement(db, e)
    if export_type == "csv":
        def safe_cell(value):
            if not isinstance(value, str) or not value: return value
            if value[0] in "=+@\t\r": return "'" + value
            if value.startswith("-"):
                try: float(value.replace(",", ""))
                except ValueError: return "'" + value
            return value
        output = io.StringIO(); writer = csv.DictWriter(output, fieldnames=columns); writer.writeheader(); writer.writerows({key: safe_cell(value) for key, value in row.items()} for row in rows)
        return Response(output.getvalue(), media_type="text/csv", headers={"Content-Disposition": f'attachment; filename="{e.public_id}-workpaper.csv"'})
    events = list_events(engagement_id, user, db) if "events" in granted else []
    return {"engagement": engagement_out(db, e, user), "artifact_id": a.public_id if a else None, "transactions": rows if "transactions" in granted else [], "events": events}

